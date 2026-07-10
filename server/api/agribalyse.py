"""Agribalyse data storage, loading and ingredient-to-row matching.

This module is the single owner of *how* Agribalyse data is acquired and stored,
so that the backing format (currently a CSV built from the official XLSX export)
can be swapped out later without touching the rest of the codebase.

The Agribalyse Synthese data holds one row per food, identified by two codes:

* the first column ``code``:        the Agribalyse code
* the second column ``ciqual_code``: the Ciqual food code

To match an ingredient, we look up its taxonomy properties (searched on the
node then all its parents) in priority order. Properties prefixed with
``agribalyse_`` match the first column, those prefixed with ``ciqual_`` match
the second column.
"""

from __future__ import annotations

import csv
import logging
import sys
import urllib.request
import zipfile
from pathlib import Path
from typing import Any, Callable, Optional

import openfoodfacts.taxonomy as taxonomy
import openpyxl

from api.settings import get_settings

logger = logging.getLogger(__name__)

# Column holding the Agribalyse code in the CSV.
AGRIBALYSE_CODE_COLUMN = "code"
# Column holding the Ciqual food code in the CSV.
CIQUAL_CODE_COLUMN = "ciqual_code"

# --- Agribalyse data source ------------------------------------------------

# Hardcoded Agribalyse data source (persistent DOI handle on the French
# research data repository).
SOURCE_URL = (
    "https://entrepot.recherche.data.gouv.fr/api/access/datafile/"
    ":persistentId?persistentId=doi:10.57745/JPY8JS"
)
# Name of the XLSX sheet holding the consolidated Synthese data.
SYNTHESE_SHEET = "Synthese"
# Row (1-indexed) of the first data line in the Synthese sheet (the header sits
# two rows above it).
FIRST_DATA_ROW = 5


# --- Schema -----------------------------------------------------------------


def _text(value: str | None) -> str:
    """Convert a CSV cell into a stripped string (``""`` for blank values)."""
    return (value or "").strip()


def _number(value: str | None) -> Optional[float]:
    """Convert a CSV cell into a float, returning ``None`` for blank values."""
    if value is None:
        return None
    value = value.strip()
    return float(value) if value else None


# Mapping of CSV column name -> conversion callable used to turn the raw CSV
# string into a typed value. The insertion order defines the column order of
# the stored file, so this dict *is* the schema of the Agribalyse storage.
HEADERS: dict[str, Callable[[str | None], Any]] = {
    "code": _text,
    "ciqual_code": _text,
    "group": _text,
    "subgroup": _text,
    "name_fr": _text,
    "lci_name": _text,
    "season_code": _text,
    "plane_code": _text,
    "delivery": _text,
    "packaging_approach": _text,
    "preparation": _text,
    "dqr": _number,
    "score": _number,
    "cc": _number,
    "ozone": _number,
    "ri": _number,
    "fpo": _number,
    "pm": _number,
    "htnhs": _number,
    "htchs": _number,
    "ated": _number,
    "eed": _number,
    "eem": _number,
    "eet": _number,
    "ecosystem_toxicity": _number,
    "land_use": _number,
    "water_use": _number,
    "energy_use": _number,
    "mineral_use": _number,
    "cc_biogenic": _number,
    "cc_fossil": _number,
    "cc_luc": _number,
}


def convert_row(raw: dict[str, str]) -> dict[str, Any]:
    """Convert a raw CSV string row into typed values using ``HEADERS``.

    Missing columns default to ``None`` so the converters handle blanks
    uniformly. Columns not part of the schema are dropped.
    """
    return {col: convert(raw.get(col)) for col, convert in HEADERS.items()}


# --- Fetch / build the CSV from the XLSX export -----------------------------


def is_xlsx(path: Path) -> bool:
    """Quick check that a file is really an XLSX (zip) archive."""
    try:
        with zipfile.ZipFile(path) as zf:
            return any(name.startswith("xl/") for name in zf.namelist())
    except (zipfile.BadZipFile, FileNotFoundError):
        return False


def download(url: str, dest: Path) -> None:
    """Download ``url`` into ``dest``."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    print(f"  downloading {url} -> {dest}", file=sys.stderr)
    with urllib.request.urlopen(url, timeout=120) as resp:
        with dest.open("wb") as fh:
            fh.write(resp.read())


def read_synthese(xlsx_path: Path) -> list[list]:
    """Read the ``Synthese`` sheet and return its data rows.

    The header (column names and order) is defined by :data:`HEADERS`, so only
    the data rows are returned. Empty rows are skipped.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if SYNTHESE_SHEET not in wb.sheetnames:
        raise KeyError(
            f"No '{SYNTHESE_SHEET}' sheet in {xlsx_path.name} "
            f"(available: {', '.join(wb.sheetnames)})"
        )
    ws = wb[SYNTHESE_SHEET]
    data: list[list] = []
    for raw in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
            continue  # skip empty rows
        data.append(list(raw))
    wb.close()
    return data


def write_csv(rows: list[list], out_path: Path) -> None:
    """Write ``rows`` to a CSV file using :data:`HEADERS` as the column header."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(list(HEADERS.keys()))
        writer.writerows(rows)
    print(f"  wrote {len(rows)} rows -> {out_path}", file=sys.stderr)


def fetch_agribalyse(out_path: Path, cache_dir: Path, no_cache: bool = False) -> int:
    """Download the Agribalyse XLSX and dump its Synthese sheet to a CSV.

    The XLSX is cached under ``cache_dir`` and re-used unless ``no_cache`` is
    set or the cached file is missing/invalid. Returns ``0`` on success,
    non-zero on error.
    """
    cache_dir.mkdir(parents=True, exist_ok=True)
    dest = cache_dir / "agribalyse.xlsx"
    if no_cache or not dest.exists() or not is_xlsx(dest):
        download(SOURCE_URL, dest)
    if not is_xlsx(dest):
        print(f"ERROR: {dest} is not a valid XLSX file", file=sys.stderr)
        return 1

    rows = read_synthese(dest)
    write_csv(rows, out_path)
    return 0


# --- Load the stored CSV into lookup indexes -------------------------------

# Lazy caches for the Agribalyse lookup tables, keyed by code column value.
_agribalyse_by_code: Optional[dict[str, dict[str, Any]]] = None
_agribalyse_by_ciqual: Optional[dict[str, dict[str, Any]]] = None


def _load_agribalyse() -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    """Load the Agribalyse CSV into two indexes (by code, by Ciqual code).

    Each raw CSV string row is converted to typed values through :data:`HEADERS`
    before being indexed. The result is cached for the lifetime of the process.
    """
    global _agribalyse_by_code, _agribalyse_by_ciqual
    if _agribalyse_by_code is not None and _agribalyse_by_ciqual is not None:
        return _agribalyse_by_code, _agribalyse_by_ciqual

    csv_path = get_settings().agribalyse_csv_path
    by_code: dict[str, dict[str, Any]] = {}
    by_ciqual: dict[str, dict[str, Any]] = {}
    with csv_path.open(newline="", encoding="utf-8") as fh:
        for raw in csv.DictReader(fh):
            row = convert_row(raw)
            code = row[AGRIBALYSE_CODE_COLUMN]
            if code:
                by_code[code] = row
            ciqual = row[CIQUAL_CODE_COLUMN]
            if ciqual:
                by_ciqual[ciqual] = row
    logger.info(
        "Loaded %d Agribalyse rows (%d ciqual) from %s",
        len(by_code),
        len(by_ciqual),
        csv_path,
    )
    _agribalyse_by_code = by_code
    _agribalyse_by_ciqual = by_ciqual
    return by_code, by_ciqual


# --- Ingredient matching ----------------------------------------------------

# Taxonomy property names searched, in priority order, on each ingredient node
# and its parents. The CSV column is derived from the property prefix.
INGREDIENT_CODE_PROPERTIES: list[str] = [
    "agribalyse_food_code",
    "agribalyse_proxy_food_code",
    "ciqual_food_code",
    "ciqual_proxy_food_code",
]


def _column_for_property(prop: str) -> str:
    """Return the CSV column used to look up a code coming from `prop`."""
    if prop.startswith("agribalyse"):
        return AGRIBALYSE_CODE_COLUMN
    if prop.startswith("ciqual"):
        return CIQUAL_CODE_COLUMN
    raise ValueError(f"Unknown code property prefix: {prop!r}")


def _node_chain(node: taxonomy.TaxonomyNode) -> list[taxonomy.TaxonomyNode]:
    """Return the node followed by all its parents (closest first)."""
    return [node, *node.get_parents_hierarchy()]


def _property_value(node: taxonomy.TaxonomyNode, prop: str) -> Optional[str]:
    """Read a code property from a taxonomy node.

    Taxonomy properties are stored as language -> value dicts (e.g.
    `{"en": "25525"}`); we return the value, preferring English then any.
    """
    raw = node.properties.get(prop)
    if raw is None:
        return None
    if isinstance(raw, dict):
        value = raw.get("en") or next(iter(raw.values()), None)
        return str(value) if value is not None else None
    return str(raw)


def find_agribalyse_row(
    node: Optional[taxonomy.TaxonomyNode],
) -> tuple[Optional[str], Optional[str], Optional[dict[str, Any]]]:
    """Find the Agribalyse row matching a taxonomy node.

    The node's code properties (and those of its parents) are tried in priority
    order. Returns ``(matched_code, code_source_property, row)`` where ``row``
    is ``None`` when nothing matched.
    """
    if node is None:
        return None, None, None

    by_code, by_ciqual = _load_agribalyse()
    indexes = {AGRIBALYSE_CODE_COLUMN: by_code, CIQUAL_CODE_COLUMN: by_ciqual}

    for prop in INGREDIENT_CODE_PROPERTIES:
        column = _column_for_property(prop)
        index = indexes[column]
        for n in _node_chain(node):
            code = _property_value(n, prop)
            if code and code in index:
                return code, prop, index[code]
    return None, None, None
